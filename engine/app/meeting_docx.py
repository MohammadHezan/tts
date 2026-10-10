"""The meeting summary as a Word document: the notes of the whole meeting, then
each speaker's summary and everything they said (Arabic as spoken, with its
English translation)."""

from __future__ import annotations

import datetime
import io
import re
from collections.abc import Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.meeting_record import Phrase
from app.summary import bullets_only, parse_notes

_ARABIC = re.compile("[؀-ۿݐ-ݿ]")
NOTE = (
    "What was said in Arabic was translated into English by the interpreter, and the summaries are written by a local "
    "AI model from that English text. They can contain mistakes: check anything important against the sentences below."
)


def _rtl(paragraph, text: str) -> None:
    """Right-to-left paragraph with its run marked right-to-left, so Word lays Arabic out properly."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    run = paragraph.add_run(text)
    run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
    run.font.name = "Arial"
    fonts = run._r.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is not None:
        fonts.set(qn("w:cs"), "Arial")


def _text(paragraph, text: str, **style) -> None:
    run = paragraph.add_run(text)
    run.italic = style.get("italic")
    run.bold = style.get("bold")
    if (color := style.get("color")) is not None:
        run.font.color.rgb = RGBColor(*color)


def _notes(doc, parts: Sequence[tuple[str, str]]) -> None:
    for kind, text in parts:
        if kind == "heading":
            doc.add_paragraph().add_run(text).bold = True
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)


def _clock(ms: int, tz: datetime.timezone) -> str:
    return datetime.datetime.fromtimestamp(ms / 1000, tz).strftime("%H:%M:%S")


def build_docx(
    phrases: Sequence[Phrase],
    speakers: Sequence[str],
    overall: str,
    per_speaker: dict[str, str],
    tz: datetime.timezone = datetime.timezone.utc,
    with_summary: bool = True,
) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Meeting summary", 0)

    if phrases:
        start, end = phrases[0].start_ms, phrases[-1].end_ms
        minutes = max(1, round((end - start) / 60000))
        when = datetime.datetime.fromtimestamp(start / 1000, tz)
        doc.add_paragraph(f"{when:%A %d %B %Y}, {_clock(start, tz)} to {_clock(end, tz)} ({minutes} min)")
    counts = {name: sum(1 for p in phrases if p.speaker == name) for name in speakers}
    doc.add_paragraph("Speakers: " + ", ".join(f"{name} ({counts[name]} phrases)" for name in speakers))
    _text(doc.add_paragraph(), NOTE, italic=True, color=(0x59, 0x59, 0x59))

    doc.add_heading("The meeting", 1)
    if with_summary:
        _notes(doc, parse_notes(overall))
    else:
        doc.add_paragraph("The meeting is still going, so only the sentences are listed. The summaries are written once it has ended.")

    doc.add_heading("By speaker", 1)
    for name in speakers:
        doc.add_heading(name, 2)
        if with_summary:
            doc.add_paragraph().add_run("Summary").bold = True
            _notes(doc, bullets_only(per_speaker.get(name, "")) or [("paragraph", "No summary could be written.")])
        doc.add_paragraph().add_run("What they said").bold = True
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for phrase in (p for p in phrases if p.speaker == name):
            time_cell, text_cell = table.add_row().cells
            time_cell.width = Pt(60)
            time_cell.text = _clock(phrase.start_ms, tz)
            first = text_cell.paragraphs[0]
            if _ARABIC.search(phrase.source):
                _rtl(first, phrase.source)
                if phrase.english:
                    _text(text_cell.add_paragraph(), phrase.english, italic=True, color=(0x1F, 0x4E, 0x79))
            else:
                first.add_run(phrase.source)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
