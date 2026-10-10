"""The meeting record, speaker matching, the summaries and the Word document, and
the dashboard endpoint - with a fake model and a fake Attendee, no GPU:

    python test_meeting_summary.py
"""

from __future__ import annotations

import asyncio
import datetime
import io
import os
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, ".")

import yaml  # noqa: E402
from docx import Document  # noqa: E402

from app.meeting_docx import build_docx  # noqa: E402
from app.meeting_record import MeetingRecorder, attribute_speakers  # noqa: E402
from app.schema import EventType, PipelineEvent  # noqa: E402
from app.summary import bullets_only, parse_notes, split_lines, summarize_meeting, summarize_person  # noqa: E402
from app.summary_service import NothingToSummarize, make_summary  # noqa: E402

FAILED: list[str] = []


def check(name: str, ok: bool, detail: object = "") -> None:
    print(("ok   " if ok else "FAIL ") + name + ("" if ok else f"  {detail}"))
    if not ok:
        FAILED.append(name)


T0 = 1_800_000_000_000  # the meeting's first moment, epoch ms


def heard(turn: str, lang: str, text: str, end_ms: int, length_ms: int = 3000) -> PipelineEvent:
    # as the pipeline emits it: stamped when produced, `latency_ms` after the speaker stopped
    return PipelineEvent(type=EventType.FINAL, lang=lang, text=text, turn_id=turn, latency_ms=500.0,
                         timestamp=(end_ms + 500) / 1000, speech_ms=float(length_ms))


def translated(turn: str, lang: str, text: str, merged: list[str] | None = None) -> PipelineEvent:
    return PipelineEvent(type=EventType.TRANSLATION, lang=lang, text=text, turn_id=turn, merged_turn_ids=merged or [])


def speech(name: str, uuid: str, start: int, stop: int) -> list[dict]:
    return [{"participant_name": name, "participant_uuid": uuid, "event_type": "speech_start", "timestamp_ms": T0 + start},
            {"participant_name": name, "participant_uuid": uuid, "event_type": "speech_stop", "timestamp_ms": T0 + stop}]


def record_sample(recorder: MeetingRecorder, bot: str = "bot_1") -> None:
    for event in (
        heard("t1", "en", "Good morning Ahmad, shall we start with the walnut tables?", T0 + 4000),
        translated("t1", "ar", "صباح الخير يا أحمد، هل نبدأ بطاولات الجوز؟"),
        heard("t2", "ar", "صباح الخير، نريد اثني عشر طاولة", T0 + 9000),
        heard("t3", "ar", "والتسليم قبل نهاية آذار", T0 + 12000),
        translated("t3", "en", "Good morning, we want twelve tables, and delivery before the end of March.", merged=["t2"]),
        heard("t4", "en", "That works, send the deposit this week.", T0 + 17000),
        translated("t4", "ar", "هذا مناسب، أرسل العربون هذا الأسبوع."),
    ):
        recorder.record(bot, event)


async def main() -> int:
    tmp = Path(tempfile.mkdtemp())
    recorder = MeetingRecorder(tmp / "meetings")
    check("record off by default", not MeetingRecorder(None).enabled)
    record_sample(recorder)
    phrases = recorder.load("bot_1")
    check("four phrases recorded, in order", [p.turn_id for p in phrases] == ["t1", "t2", "t3", "t4"], [p.turn_id for p in phrases])
    check("English phrases are their own English", phrases[0].english.startswith("Good morning Ahmad") and phrases[3].english.startswith("That works"))
    check("a merged batch's translation sits on its last phrase", phrases[1].english == "" and phrases[2].english.startswith("Good morning, we want twelve"), [p.english for p in phrases])
    check("times come from the event", phrases[0].end_ms == T0 + 4000 and phrases[0].start_ms == T0 + 1000, (phrases[0].start_ms, phrases[0].end_ms))
    check("the file is on disk", (tmp / "meetings" / "bot_1.jsonl").exists())
    check("an unknown meeting has no phrases", recorder.load("nope") == [])
    with (tmp / "meetings" / "bot_1.jsonl").open("a", encoding="utf-8") as f:
        f.write('{"kind": "heard", "turn"')  # a half-written last line
    check("a half-written last line is ignored", len(recorder.load("bot_1")) == 4)

    # who spoke
    events = speech("Sarah", "u-sarah", 800, 4300) + speech("Ahmad", "u-ahmad", 6000, 12500) + speech("Sarah", "u-sarah", 14000, 17400)
    order = attribute_speakers(phrases, events)
    check("speakers matched by who was speaking", [p.speaker for p in phrases] == ["Sarah", "Ahmad", "Ahmad", "Sarah"], [p.speaker for p in phrases])
    check("speakers in order of first speech", order == ["Sarah", "Ahmad"], order)
    lone = [p for p in recorder.load("bot_1")]
    attribute_speakers(lone, speech("Ahmad", "u-1", 10_000_000, 10_001_000))
    check("nobody was speaking near it: unknown speaker", all(p.speaker == "Unknown speaker" for p in lone))
    close = recorder.load("bot_1")
    attribute_speakers(close, speech("Sarah", "u-sarah", 5200, 6000))  # spoke 1.2s after t1 ended: close enough
    check("nearest speaker within 2.5s is used", close[0].speaker == "Sarah", close[0].speaker)
    far = recorder.load("bot_1")
    attribute_speakers(far, speech("Sarah", "u-sarah", 8000, 8500))  # 4s after t1 ended: too far
    check("a speaker further than 2.5s away is not used", far[0].speaker == "Unknown speaker", far[0].speaker)
    twins = recorder.load("bot_1")
    attribute_speakers(twins, speech("Sam", "u-a", 800, 4300) + speech("Sam", "u-b", 6000, 12500))
    check("two people with the same name stay apart", [p.speaker for p in twins[:3]] == ["Sam", "Sam (2)", "Sam (2)"], [p.speaker for p in twins])
    nameless = recorder.load("bot_1")
    attribute_speakers(nameless, [{**e, "participant_name": ""} for e in speech("", "u-x", 800, 4300)])
    check("someone without a name is 'Speaker 1'", nameless[0].speaker == "Speaker 1", nameless[0].speaker)

    # the model's text
    check("lines are split between sentences", split_lines(["aaa", "bbb", "ccc"], 7) == ["aaa\nbbb", "ccc"] and split_lines([], 5) == [])
    parts = parse_notes("Overview: A short one.\nKey points:\n- one\n* two\n**Decisions and action items:**\n- Ahmad will pay.")
    check("notes are split into headings, text and bullets", parts == [("heading", "Overview"), ("paragraph", "A short one."), ("heading", "Key points"),
                                                                        ("bullet", "one"), ("bullet", "two"), ("heading", "Decisions and action items"), ("bullet", "Ahmad will pay.")], parts)
    check("a speaker's summary drops the lead-in line", bullets_only("Here is a summary:\n- he agreed\n- he pays") == [("bullet", "he agreed"), ("bullet", "he pays")])

    calls: list[str] = []

    async def fake_ask(prompt: str) -> str:
        calls.append(prompt)
        if "Summarize what" in prompt:
            return "Here you go:\n- said something useful"
        return "Overview: They agreed on twelve tables.\nKey points:\n- Twelve walnut tables\nDecisions and action items:\n- Ahmad pays the deposit this week."

    await summarize_meeting(fake_ask, ["Sarah: hello"] * 3)
    check("a short meeting takes one model call", len(calls) == 1, len(calls))
    calls.clear()
    await summarize_meeting(fake_ask, [f"Sarah: {'word ' * 100}"] * 40, limit=3000)
    check("a long meeting goes in pieces, then is combined", len(calls) >= 4 and "consecutive parts" in calls[-1], len(calls))
    calls.clear()
    await summarize_person(fake_ask, "Ahmad", [f"Ahmad: {'word ' * 100}"] * 20, limit=3000)
    check("a long speaker goes in pieces too", len(calls) >= 3 and "Ahmad" in calls[-1], len(calls))

    # the document
    calls.clear()
    attribute_speakers(phrases, events)
    data = await make_summary(recorder, "bot_1", lambda: _events(events), fake_ask, datetime.timezone(datetime.timedelta(hours=3)))
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    check("the document opens and has the meeting notes", "They agreed on twelve tables." in text and "Ahmad pays the deposit" in text)
    check("it has a section per speaker, with a summary", "Sarah" in text and "Ahmad" in text and text.count("said something useful") == 2)
    check("it lists what they said", any("walnut tables" in c for c in cells) and any("That works" in c for c in cells))
    check("Arabic is there as spoken, with its English", any("اثني عشر" in c for c in cells) or any("والتسليم" in c for c in cells))
    check("and the English translation of it", any("we want twelve tables" in c for c in cells))
    xml = doc.element.xml
    check("Arabic is laid out right-to-left", "<w:bidi/>" in xml and "<w:rtl/>" in xml)
    check("times use the zone given (UTC+3)", f"{datetime.datetime.fromtimestamp((T0 + 1000) / 1000, datetime.timezone(datetime.timedelta(hours=3))):%H:%M:%S}" in " ".join(cells))

    # during a meeting: sentences only, no model
    calls.clear()
    data = await make_summary(recorder, "bot_1", lambda: _events(events), fake_ask, with_summary=False)
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    check("during the meeting no model is used", calls == [], calls)
    check("and the document says so", "still going" in text and "said something useful" not in text)

    async def broken() -> list:
        raise OSError("Attendee is down")

    data = await make_summary(recorder, "bot_1", broken, fake_ask)
    check("without Attendee's speaker data it still works", "Unknown speaker" in "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs))
    try:
        await make_summary(recorder, "nobody", lambda: _events([]), fake_ask)
        check("nothing recorded: nothing to summarize", False)
    except NothingToSummarize:
        check("nothing recorded: nothing to summarize", True)

    await endpoint(tmp)
    print(f"\n{'ALL PASSED' if not FAILED else 'FAILED: ' + ', '.join(FAILED)}")
    return 1 if FAILED else 0


async def _events(events: list[dict]) -> list[dict]:
    return events


async def endpoint(tmp: Path) -> None:
    """The real server: live meeting -> sentences only; ended -> summary written once and kept."""
    config = yaml.safe_load(Path("../deploy/config.docker-gpu.yaml").read_text(encoding="utf-8"))
    config["tts"]["provider"] = "none"
    config["auth"] = {"username": "x", "password_hash": ""}  # login is tested in test_auth.py
    config["meeting"]["record_dir"] = str(tmp / "served")
    path = tmp / "config.yaml"
    path.write_text(yaml.safe_dump(config, allow_unicode=True), encoding="utf-8")
    os.environ["ENGINE_CONFIG_PATH"] = str(path)
    os.environ["INTERPRETER_ADMIN_PASSWORD_HASH"] = ""  # login off here, whatever the .env of this computer says
    from starlette.testclient import TestClient

    from app import server

    recorder = server._recorder
    record_sample(recorder, "bot_9")
    asked: list[str] = []

    def fake_asker(*_args, **_kwargs):
        async def ask(prompt: str) -> str:
            asked.append(prompt)
            return "Overview: Done.\n- a point" if "Overview" in prompt else "- did something"
        return ask

    async def fake_events(client, bot_id):
        return speech("Sarah", "u-s", 800, 4300) + speech("Ahmad", "u-a", 6000, 12500) + speech("Sarah", "u-s", 14000, 17400)

    server.ollama_asker, server.fetch_participant_events = fake_asker, fake_events
    client = TestClient(server.app)
    server._bot_controls.heard_from("bot_9")  # its audio is arriving: the meeting is going on
    live = client.get("/api/bots/bot_9/summary.docx")
    check("live meeting: a document comes at once", live.status_code == 200 and "sentences-so-far" in live.headers["content-disposition"], live.status_code)
    check("live meeting: the model is not asked", asked == [], asked)
    check("live meeting: nothing is saved as the summary", not recorder.summary_path("bot_9").exists())

    server._bot_controls._heard["bot_9"] = time.monotonic() - 120  # audio stopped two minutes ago
    done = client.get("/api/bots/bot_9/summary.docx")
    check("after the meeting: the summary is written", done.status_code == 200 and asked and "Done." in "\n".join(p.text for p in Document(io.BytesIO(done.content)).paragraphs), done.status_code)
    check("and kept on this PC", recorder.summary_path("bot_9").exists())
    asked.clear()
    again = client.get("/api/bots/bot_9/summary.docx")
    check("the next download is the saved one: no model run", again.status_code == 200 and asked == [] and again.content == done.content)
    client.get("/api/bots/bot_9/summary.docx?refresh=true")
    check("refresh writes it again", bool(asked))
    check("no record for this meeting: 404", client.get("/api/bots/zzz/summary.docx").status_code == 404)

    # the automatic step: the bridge disconnects, Attendee says the meeting ended, the summary is written
    record_sample(recorder, "bot_10")
    asked.clear()
    server._bot_controls._heard["bot_10"] = time.monotonic() - 120
    states = iter(["joined_recording", "ended"])

    async def attendee_state(call):
        return {"state": next(states)}

    real_sleep, real_call = asyncio.sleep, server._call_attendee

    async def instant(_seconds, *args, **kwargs):
        return await real_sleep(0)

    server._call_attendee, asyncio.sleep = attendee_state, instant
    try:
        await server._write_summary_when_over("bot_10")
    finally:
        server._call_attendee, asyncio.sleep = real_call, real_sleep
    check("the summary is written by itself once the meeting has ended", recorder.summary_path("bot_10").exists() and bool(asked))

    record_sample(recorder, "bot_11")
    asked.clear()
    server._bot_controls.heard_from("bot_11")  # its audio is back: it was only a dropped connection
    await server._write_summary_when_over("bot_11")
    check("a dropped connection that comes back is not the end of the meeting", not recorder.summary_path("bot_11").exists() and asked == [])


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
